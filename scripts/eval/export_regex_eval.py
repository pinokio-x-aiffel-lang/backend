"""562개 클레임 전체를 Excel로 출력 — regex 결과 + GT 포함.

Usage:
    uv run python scripts/eval/export_regex_eval.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).parents[2]))

from src.modules.normalize_claim import _parse_value

JSONL = Path("data/eval/labeling_base_success.jsonl")
OUT   = Path("data/eval/regex_eval_562.xlsx")
SKIP  = {"불명", ""}

# 색상
GREEN  = PatternFill("solid", fgColor="C6EFCE")
RED    = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
HEADER = PatternFill("solid", fgColor="4472C4")


def main() -> None:
    claims = []
    with JSONL.open(encoding="utf-8") as f:
        for line in f:
            obj = json.loads(line)
            for c in obj["claims"]:
                c["_title"] = obj["title"]
                c["_published_at"] = obj["published_at"]
            claims.extend(obj["claims"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "regex_eval"

    # 헤더
    headers = [
        "no", "article_id", "published_at", "title",
        "claim_id", "sentence",
        "value_raw", "GT (llm_value)", "regex 결과", "일치",
        "period_raw", "GT (period)", "regex period",
    ]
    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # 열 너비
    widths = [5, 28, 12, 35, 28, 55, 25, 20, 20, 8, 20, 15, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    match_v = mismatch_v = 0

    for no, c in enumerate(claims, 1):
        v_raw  = c["value"]["raw"]
        v_gt   = c["value"]["llm_value"]
        v_got  = _parse_value(v_raw)

        p_raw  = c["period_value"]["raw"]
        p_gt   = c["period_value"]["llm_value"]
        from src.modules.normalize_claim import _normalize_period
        p_got  = _normalize_period(p_raw)

        v_match = (v_got == v_gt)
        if v_raw not in SKIP and v_gt not in SKIP:
            if v_match:
                match_v += 1
            else:
                mismatch_v += 1

        row = [
            no,
            c["article_id"],
            c["_published_at"],
            c["_title"],
            c["claim_id"],
            c["sentence"],
            v_raw,
            v_gt,
            v_got,
            "O" if v_match else "X",
            p_raw,
            p_gt,
            p_got,
        ]
        ws.append(row)

        # 행 색상
        r = no + 1
        fill = GREEN if v_match else RED
        for col in range(7, 11):   # value 관련 열
            ws.cell(r, col).fill = fill

        # 셀 줄바꿈
        for col in [6]:
            ws.cell(r, col).alignment = Alignment(wrap_text=True)

    # 요약 시트
    ws2 = wb.create_sheet("요약")
    total = match_v + mismatch_v
    ws2.append(["항목", "값"])
    ws2.append(["전체 클레임", len(claims)])
    ws2.append(["평가 대상 (raw·GT 모두 불명 아닌 것)", total])
    ws2.append(["MATCH (일치)", match_v])
    ws2.append(["MISMATCH (불일치)", mismatch_v])
    ws2.append(["일치율", f"{match_v/total*100:.1f}%" if total else "-"])

    wb.save(OUT)
    print(f"저장 완료: {OUT}")
    print(f"MATCH {match_v}/{total} ({match_v/total*100:.1f}%)")


if __name__ == "__main__":
    main()
