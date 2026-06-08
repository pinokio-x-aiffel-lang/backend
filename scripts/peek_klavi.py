import openpyxl

wb = openpyxl.load_workbook(
    r'D:\Cursor\Project\AIFFEL thon\Data\클라비 제공 뉴스 기사 데이터.xlsx',
    read_only=True, data_only=True
)
print('시트:', wb.sheetnames)
ws = wb.active
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
print('컬럼:', headers)
print('총 행수:', ws.max_row)
print()
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=4, values_only=True), 1):
    print(f'[{i}]', row[:5])
